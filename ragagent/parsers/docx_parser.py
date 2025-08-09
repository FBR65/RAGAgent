from pathlib import Path
from typing import List
import logging
from .base import BaseParser
from ..models import DocumentType

logger = logging.getLogger(__name__)


class DOCXParser(BaseParser):
    """Parser for Word documents"""

    def can_parse(self, file_path: Path) -> bool:
        """Check if file is a DOCX file"""
        return file_path.suffix.lower() == ".docx"

    def parse(self, file_path: Path) -> str:
        """Extract text from DOCX file"""
        try:
            from docx import Document
        except ImportError:
            raise ImportError(
                "python-docx is required for DOCX parsing. Install with: pip install python-docx"
            )

        try:
            doc = Document(file_path)
            text = ""

            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\t"
                text += "\n"

            if not text.strip():
                logger.warning(f"No text extracted from DOCX: {file_path}")
                return ""

            logger.info(f"Extracted {len(text)} characters from DOCX: {file_path}")
            return text.strip()

        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path}: {e}")
            raise

    def get_supported_extensions(self) -> List[str]:
        """Return supported DOCX extensions"""
        return [".docx"]
